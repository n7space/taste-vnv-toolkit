"""Extract requirements from Excel file to Doxygen tag file."""

from vnvtoolkit import get_setting
import os
import xml.etree.ElementTree as ET
from xml.dom import minidom

ISO_LATIN_LETTER_COUNT = 26

def emit_progress(value):
    """Report progress if the function is available."""
    if "report_progress" in globals():
        report_progress(value)


def column_letter_to_index(col_letter):
    """Convert Excel column letter (A, B, C, ..., Z, AA, AB, ...) to 0-based index."""
    col_letter = col_letter.upper().strip()
    index = 0
    for char in col_letter:
        index = index * ISO_LATIN_LETTER_COUNT + (ord(char) - ord('A') + 1)
    return index - 1


def read_excel_requirements(excel_path, sheet_name, id_col, title_col, data_start_row):
    """
    Read requirements from Excel file.
    Returns list of tuples: [(req_id, req_title), ...]
    """
    try:
        import openpyxl
    except ImportError:
        raise Exception("openpyxl library is required. Install with: pip install openpyxl")
    
    id_col_idx = column_letter_to_index(id_col)
    title_col_idx = column_letter_to_index(title_col)
    
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise Exception(f"Sheet '{sheet_name}' not found in Excel file. Available sheets: {', '.join(wb.sheetnames)}")
    
    ws = wb[sheet_name]
    requirements = []
    
    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        if len(row) <= max(id_col_idx, title_col_idx):
            continue
        
        req_id = row[id_col_idx]
        req_title = row[title_col_idx]
        
        # Skip rows with empty ID or title
        if not req_id or not req_title:
            continue
        
        # Convert to string and strip whitespace
        req_id = str(req_id).strip()
        req_title = str(req_title).strip()
        
        # Skip empty strings after stripping
        if not req_id or not req_title:
            continue
        
        requirements.append((req_id, req_title))
    
    wb.close()
    return requirements


def create_doxygen_tag_file(requirements, output_path):
    """
    Create Doxygen tag file from requirements list.
    requirements: list of tuples [(req_id, req_title), ...]
    """
    # Create root element
    root = ET.Element('tagfile')
    root.set('doxygen_version', '1.16.0')
    
    # Add each requirement as a compound element
    for req_id, req_title in requirements:
        compound = ET.SubElement(root, 'compound')
        compound.set('kind', 'requirement')
        
        id_elem = ET.SubElement(compound, 'id')
        id_elem.text = req_id
        
        title_elem = ET.SubElement(compound, 'title')
        title_elem.text = req_title
    
    # Pretty print the XML
    xml_str = ET.tostring(root, encoding='utf-8')
    dom = minidom.parseString(xml_str)
    pretty_xml = dom.toprettyxml(indent='  ', encoding='UTF-8')
    
    # Write to file
    with open(output_path, 'wb') as f:
        f.write(pretty_xml)


# ── Main tool execution ───────────────────────────────────────────────────────

if not taste_project_directory:
    status = "error"
    status_text = "Project directory is not configured"
    show_status = True
else:
    try:
        emit_progress(5)
        
        # Get settings
        excel_file = str(get_setting(settings, "Excel file name", "demo_requirements.xlsx"))
        sheet_name = str(get_setting(settings, "Sheet name", "5. Requirements"))
        id_col = str(get_setting(settings, "ID column", "A"))
        title_col = str(get_setting(settings, "Title column", "C"))
        data_start_row = int(get_setting(settings, "Data start row", 3))
        output_tag_file = str(get_setting(settings, "Output tag file", "doxygen-requirements.tag"))
        
        emit_progress(10)
        
        # Build paths
        excel_path = os.path.join(taste_project_directory, excel_file)
        output_path = os.path.join(taste_project_directory, output_tag_file)
        
        # Check if Excel file exists
        if not os.path.isfile(excel_path):
            status = "error"
            status_text = f"Excel file not found: {excel_path}"
            show_status = True
        else:
            emit_progress(20)
            
            # Read requirements from Excel
            requirements = read_excel_requirements(
                excel_path, 
                sheet_name, 
                id_col, 
                title_col, 
                data_start_row
            )
            
            emit_progress(60)
            
            if not requirements:
                status = "warning"
                status_text = f"No requirements found in {excel_file}, sheet '{sheet_name}'"
                show_status = True
            else:
                # Create Doxygen tag file
                create_doxygen_tag_file(requirements, output_path)
                
                emit_progress(100)
                
                status = "ok"
                status_text = f"Successfully extracted {len(requirements)} requirement(s) from {excel_file}\nCreated tag file: {output_path}"
                show_status = True
                
    except Exception as exc:
        status = "error"
        status_text = f"Failed to extract requirements: {exc}"
        show_status = True
